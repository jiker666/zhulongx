#!/usr/bin/env python3
"""Build the formal Word versions of the competition delivery documents."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "docs" / "competition"
OUTPUT_DIR = SOURCE_DIR / "docx"
ASSET = SOURCE_DIR / "assets" / "audit-console-overview.jpg"

DOCUMENTS = [
    ("README.md", "00-delivery-index.docx", "竞赛交付索引"),
    ("01-solution-design.md", "01-solution-design.docx", "方案设计文档"),
    ("02-development-guide.md", "02-development-guide.docx", "开发文档"),
    ("03-test-report.md", "03-test-report.docx", "测试报告"),
    ("04-user-and-deployment-guide.md", "04-user-and-deployment-guide.docx", "用户与部署手册"),
    ("05-technical-report.md", "05-technical-report.docx", "技术报告"),
]

PROJECT_LABEL = "烛龙·玄枢 | XH-202609 竞赛提交材料"
CJK_FONT = "Arial Unicode MS"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "5D6B7A"
TABLE_WIDTH = 9360


def set_run_font(run, name: str = "Calibri", size: float | None = None, color: str | None = None, bold: bool | None = None) -> None:
    latin_font = CJK_FONT if name == "Calibri" else name
    run.font.name = latin_font
    run._element.rPr.rFonts.set(qn("w:ascii"), latin_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin_font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    run._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    lang = run._element.rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.rPr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for column, width in zip(grid.gridCol_lst, widths):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_border(paragraph, color: str = "B7C4D4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def clean_text(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"`([^`]+)`", r"\1", value)
    value = value.replace("**", "").replace("__", "")
    value = value.replace("\\|", "|")
    return value.strip()


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = CJK_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = CJK_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = doc.styles.add_style("Document Title", WD_STYLE_TYPE.PARAGRAPH)
    title.font.name = CJK_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    title.font.size = Pt(24)
    title.font.color.rgb = RGBColor.from_string(INK)
    title.font.bold = True
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True

    subtitle = doc.styles.add_style("Document Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = CJK_FONT
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    subtitle.font.size = Pt(10)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(18)

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Courier New"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    code.font.size = Pt(8.5)
    code.paragraph_format.space_before = Pt(5)
    code.paragraph_format.space_after = Pt(5)
    code.paragraph_format.left_indent = Inches(0.16)
    code.paragraph_format.right_indent = Inches(0.16)

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.167


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(PROJECT_LABEL)
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    set_paragraph_border(paragraph)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("XH-202609  |  第 ")
    set_run_font(run, size=8, color=MUTED)
    add_page_field(paragraph)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=8, color=MUTED)


def add_title_block(doc: Document, title: str) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(6)
    run = kicker.add_run("XH-202609 竞赛提交材料")
    set_run_font(run, size=9, color=BLUE, bold=True)
    paragraph = doc.add_paragraph(style="Document Title")
    paragraph.add_run(title)
    subtitle = doc.add_paragraph(style="Document Subtitle")
    subtitle.add_run("烛龙·玄枢（ZhuLong XuanShu）——具备自主决策能力的通用网络安全智能体")
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(14)
    set_paragraph_border(rule, BLUE)


def add_figure(doc: Document) -> None:
    if not ASSET.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = paragraph.add_run().add_picture(str(ASSET), width=Inches(6.35))
    picture._inline.docPr.set("descr", "烛龙玄枢控制台总览界面，已脱敏，仅用于说明用户操作界面。")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(12)
    run = caption.add_run("图 1  控制台总览（脱敏示意）")
    set_run_font(run, size=9, color=MUTED)


def add_callout(doc: Document, text: str, label: str = "说明") -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(5)
    set_paragraph_shading(paragraph, LIGHT_GRAY)
    label_run = paragraph.add_run(f"{label}  ")
    set_run_font(label_run, size=10.5, color=DARK_BLUE, bold=True)
    body_run = paragraph.add_run(clean_text(text))
    set_run_font(body_run, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: list[str], language: str, compact: bool = False) -> None:
    if language == "mermaid":
        add_callout(
            doc,
            "任务与附件 → 输入归一化 → 计划与假设 → PolicyEngine → 受控工具与隔离执行 → 证据与 checkpoint → 结构化结论与审计。",
            "流程示意",
        )
        return
    paragraph = doc.add_paragraph()
    paragraph.style = "Code Block"
    paragraph.paragraph_format.left_indent = Inches(0.16)
    paragraph.paragraph_format.right_indent = Inches(0.16)
    paragraph.paragraph_format.space_before = Pt(2 if compact else 5)
    paragraph.paragraph_format.space_after = Pt(2 if compact else 0)
    set_paragraph_shading(paragraph, "F7F8FA")
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, name="Courier New", size=8.0 if compact else 8.5, color="263238")
    if not compact:
        doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table_row(line: str) -> list[str]:
    values = [clean_text(value) for value in line.strip().strip("|").split("|")]
    return values


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*", line))


def table_widths(rows: list[list[str]]) -> list[int]:
    columns = len(rows[0])
    max_lengths = [max(5, *(len(row[index]) if index < len(row) else 0 for row in rows)) for index in range(columns)]
    weights = [min(28, max(7, length)) for length in max_lengths]
    total = sum(weights)
    widths = [max(900, int(TABLE_WIDTH * weight / total)) for weight in weights]
    widths[-1] += TABLE_WIDTH - sum(widths)
    return widths


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    widths = table_widths(rows)
    set_table_geometry(table, widths)
    mark_header_row(table.rows[0])
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            size = 9.2 if len(rows[0]) >= 4 else 9.8
            set_run_font(run, size=size, color=INK, bold=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.widow_control = True
    run = paragraph.add_run(clean_text(text))
    set_run_font(run, size=11, color="1F2933")


def build_document(source: Path, output: Path, display_name: str) -> None:
    doc = Document()
    style_document(doc)
    add_header_footer(doc)
    add_title_block(doc, display_name)

    if source.name in {"04-user-and-deployment-guide.md", "05-technical-report.md"}:
        add_figure(doc)

    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    first_heading_skipped = False
    while index < len(lines):
        line = lines[index]
        if line.startswith("# ") and not first_heading_skipped:
            first_heading_skipped = True
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            add_code_block(doc, block, language, compact=source.name == "03-test-report.md")
            index += 1
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1)) - 1
            style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
            paragraph = doc.add_paragraph(style=style)
            heading_text = clean_text(heading.group(2))
            run = paragraph.add_run(heading_text)
            if source.name == "03-test-report.md" and heading_text == "9. 验收结论":
                # Keep the final report conclusion with the preceding review checklist.
                paragraph.paragraph_format.space_before = Pt(6)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.keep_with_next = True
                set_run_font(run, size=13, color=BLUE, bold=True)
            index += 1
            continue
        if line.strip().startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = [parse_table_row(table_line) for table_line in table_lines if not is_separator(table_line)]
            add_table(doc, rows)
            continue
        number = re.match(r"^\d+\.\s+(.+)$", line)
        if number:
            add_paragraph(doc, number.group(1), "List Number")
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            add_paragraph(doc, bullet.group(1), "List Bullet")
            index += 1
            continue
        if line.startswith("> "):
            add_callout(doc, line[2:], "提示")
            index += 1
            continue
        if line.strip():
            add_paragraph(doc, line)
        index += 1

    doc.core_properties.title = display_name
    doc.core_properties.subject = "XH-202609 竞赛提交材料"
    doc.core_properties.author = "ZhuLong XuanShu Team"
    doc.core_properties.comments = "由仓库 Markdown 源文档生成；不含模型密钥、平台令牌、Cookie、Authorization 或隐藏推理。"
    doc.core_properties.created = datetime(2026, 8, 31)
    doc.core_properties.modified = datetime(2026, 8, 31)
    doc.save(output)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for source_name, output_name, title in DOCUMENTS:
        build_document(SOURCE_DIR / source_name, OUTPUT_DIR / output_name, title)
        print(OUTPUT_DIR / output_name)


if __name__ == "__main__":
    main()
